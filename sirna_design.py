#!/usr/bin/env python3
"""
siRNA 设计工具 v1.0
基于 2025.7.6 siRNA设计过程 文档流程
========================================
用法:
  python sirna_design.py --gene SUN1 --species "Homo sapiens"
  python sirna_design.py --cds CDS序列文件.fasta
  python sirna_design.py --transcript NM_001130965.3
"""

import re, sys, os, time, json, argparse
import urllib.request, urllib.parse
from datetime import datetime

# ─────────────────────────────────────────────
# 配置
# ─────────────────────────────────────────────
DSIR_URL = "http://biodev.extra.cea.fr/DSIR/DSIR.php"
BLAST_URL = "https://blast.ncbi.nlm.nih.gov/blast/Blast.cgi"
NCBI_EUTILS = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils"


# ─────────────────────────────────────────────
# 第1步: 获取CDS
# ─────────────────────────────────────────────
def ncbi_fetch(url, max_retries=3):
    """NCBI请求带SSL重试"""
    for attempt in range(max_retries):
        try:
            with urllib.request.urlopen(url, timeout=45) as resp:
                return resp.read()
        except Exception as e:
            if attempt < max_retries - 1:
                time.sleep(2 ** attempt)
                continue
            raise


def fetch_cds_from_ncbi(transcript_id):
    """通过转录本ID从NCBI获取CDS序列"""
    url = f"{NCBI_EUTILS}/efetch.fcgi?db=nucleotide&id={transcript_id}&rettype=fasta_cds_na&retmode=text"
    text = ncbi_fetch(urllib.request.Request(url)).decode()

    lines = text.strip().split('\n')
    seq = ''.join(line.strip() for line in lines if not line.startswith('>'))
    header = lines[0] if lines[0].startswith('>') else '>unknown'

    if not seq or not seq.startswith('ATG'):
        raise ValueError(f"未能获取有效的CDS序列 (transcript: {transcript_id})")

    return seq, header


def search_transcript(gene_name, species="Homo sapiens"):
    """搜索基因的转录本ID，返回(tid, cds_seq, cds_header)"""
    url = f"{NCBI_EUTILS}/esearch.fcgi?db=gene&term={urllib.parse.quote(f'{gene_name}[Gene Name] AND {species}[Organism]')}&retmode=json"
    data = json.loads(ncbi_fetch(urllib.request.Request(url)).decode())

    id_list = data.get('esearchresult', {}).get('idlist', [])
    if not id_list:
        raise ValueError(f"未找到基因: {gene_name} ({species})")

    gene_id = id_list[0]
    print(f"  Gene ID: {gene_id}")

    # 获取基因XML信息提取NM转录本
    xml_text = ncbi_fetch(urllib.request.Request(
        f"{NCBI_EUTILS}/efetch.fcgi?db=gene&id={gene_id}&retmode=xml")).decode()
    nms = sorted(set(re.findall(r'NM_\d+', xml_text)))
    if not nms:
        raise ValueError(f"未找到基因{gene_name}的RefSeq转录本")

    print(f"  找到{len(nms)}个转录本: {', '.join(nms[:8])}{'...' if len(nms)>8 else ''}")

    # 选最长CDS的转录本
    best_nm = None; best_seq = None; best_header = None; best_len = 0
    for nm in nms[:20]:
        try:
            seq, header = fetch_cds_from_ncbi(nm)
            if len(seq) > best_len:
                best_len = len(seq)
                best_nm = nm; best_seq = seq; best_header = header
        except:
            continue

    if not best_nm:
        raise ValueError(f"未找到基因{gene_name}的有效CDS")
    print(f"  选择: {best_nm} (CDS: {best_len}bp, {best_len//3}aa)")
    return best_nm, best_seq, best_header


def read_fasta(filepath):
    """读取FASTA文件"""
    with open(filepath) as f:
        lines = f.readlines()
    seq = ''.join(line.strip() for line in lines if not line.startswith('>'))
    header = [l.strip() for l in lines if l.startswith('>')]
    return seq.upper(), header[0] if header else '>unknown'


# ─────────────────────────────────────────────
# 第2步: DSIR 提交与解析
# ─────────────────────────────────────────────
def submit_dsir(cds_seq, name="siRNA_design", threshold=90):
    """提交CDS到DSIR，返回所有候选siRNA"""
    fasta = f">{name}\n{cds_seq}\n"
    boundary = '----Form' + str(int(time.time()*1000))

    body_parts = []
    params = {
        'seqname': name, 'sequence_all': fasta, 'design': '19',
        'threshold': str(threshold), 'fornt': '1', 'imotifs': '1',
    }
    for key, val in params.items():
        body_parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{key}"\r\n\r\n{val}\r\n'.encode())
    body_parts.append(f'--{boundary}--\r\n'.encode())
    body_data = b''.join(body_parts)

    req = urllib.request.Request(DSIR_URL, data=body_data)
    req.add_header('Content-Type', f'multipart/form-data; boundary={boundary}')

    resp = urllib.request.urlopen(req, timeout=120)
    html = resp.read().decode('utf-8', errors='replace')

    # 解析结果表格
    rows = re.findall(r'<tr[^>]*>.*?</tr>', html, re.DOTALL)
    results = []
    for row in rows:
        vals = re.findall(r'>\s*([^<]+)\s*<', row)
        # DSIR数据行: [id, pos, ss_seq, as_seq, score, corrected]
        # checkbox列被跳过因为<input>标签内含<>导致正则不捕获
        if len(vals) >= 6 and vals[0].isdigit() and vals[1].isdigit():
            results.append({
                'id': int(vals[0]), 'pos': int(vals[1]),
                'ss': vals[2], 'as': vals[3],
                'score': float(vals[4]), 'corrected': float(vals[5]),
            })
    return results


# ─────────────────────────────────────────────
# 第3步: 筛选
# ─────────────────────────────────────────────
def gc_count(seq):
    return seq.count('G') + seq.count('C')


def has_bad_patterns(seq):
    s = seq.upper()
    # 连续2+个GC/CG
    for i in range(len(s)-3):
        if s[i:i+2] in ('GC', 'CG') and s[i+2:i+4] in ('GC', 'CG'):
            return True
    # 连续4+相同碱基
    for base in ('A', 'U', 'G', 'C'):
        if base * 4 in s:
            return True
    # UAUAUA / AUAUAU
    if 'UAUAUA' in s or 'AUAUAU' in s:
        return True
    return False


def filter_candidates(dsir_results, start_aa=100, gc_min=8, gc_max=11, verbose=True):
    """
    按文档标准筛选:
    - 位置 > start_aa
    - 第2位 = A
    - GC含量 gc_min-gc_max
    - 无bad patterns
    - F1=G/C (尽量), R1=A/U (尽量)
    """
    passed = []
    for r in dsir_results:
        pos = r['pos']; ss = r['ss']
        aa = pos // 3

        # 位置
        if pos < start_aa * 3:
            if verbose: print(f"    淘汰: pos={pos}({aa}aa) < {start_aa}aa")
            continue

        # 第2位
        if ss[1] != 'A':
            if verbose: print(f"    淘汰: pos={pos}, 第2位={ss[1]} ≠ A")
            continue

        # GC
        gc = gc_count(ss)
        if gc < gc_min or gc > gc_max:
            if verbose: print(f"    淘汰: pos={pos}, GC={gc}/19 不在{gc_min}-{gc_max}")
            continue

        # bad patterns
        if has_bad_patterns(ss):
            if verbose: print(f"    淘汰: pos={pos}, 包含bad patterns")
            continue

        # F1, R1 记录但不淘汰（文档说"尽量"）
        comp = {'A': 'U', 'U': 'A', 'G': 'C', 'C': 'G'}
        ant = ''.join(comp[b] for b in reversed(ss))
        f1_ok = ss[0] in 'GC'
        r1_ok = ant[0] in 'AU'

        passed.append({
            'pos': pos, 'aa': aa,
            'ss': ss, 'as': ant,
            'score': r['score'], 'corrected': r['corrected'],
            'gc': gc, 'gc_pct': gc/19*100,
            'f1': ss[0], 'f1_ok': f1_ok,
            'r1': ant[0], 'r1_ok': r1_ok,
        })

    return passed


def pick_best_two(candidates, min_spacing=25):
    """按3'端优先+评分+F1/R1选取两条"""
    if len(candidates) < 2:
        return candidates

    # 按位置降序（越靠近3'越好），其次F1/R1合规性，再按校正评分降序
    candidates.sort(key=lambda x: (
        -x['pos'],
        -(x['f1_ok'] + x['r1_ok']),
        -x['corrected'],
    ))

    selected = []
    for c in candidates:
        if len(selected) >= 2:
            break
        if not any(abs(c['pos'] - s['pos']) < min_spacing for s in selected):
            selected.append(c)

    return selected


# ─────────────────────────────────────────────
# 第4步: BLAST
# ─────────────────────────────────────────────
def blast_sequence(seq, gene_name=None, name="siRNA", database="refseq_rna", timeout_total=180):
    """BLAST一条序列，返回命中统计（限制在人类转录组，E-value<1为显著匹配）"""
    data = {
        'CMD': 'Put', 'PROGRAM': 'blastn', 'DATABASE': database,
        'QUERY': f">{name}\n{seq.replace('U', 'T')}",
        'ENTREZ_QUERY': 'Homo sapiens[Organism]',
        'HITLIST_SIZE': '30', 'EXPECT': '10000',
        'WORD_SIZE': '7', 'FILTER': 'F', 'FORMAT_TYPE': 'XML',
    }
    encoded = urllib.parse.urlencode(data).encode()
    req = urllib.request.Request(BLAST_URL, data=encoded)

    try:
        resp = urllib.request.urlopen(req, timeout=30)
        text = resp.read().decode()
    except Exception as e:
        return {'error': str(e), 'total': 0, 'target': 0, 'offtarget': 0}

    m = re.search(r'RID\s*=\s*(\S+)', text)
    if not m:
        return {'error': 'No RID', 'total': 0, 'target': 0, 'offtarget': 0}
    rid = m.group(1)
    print(f"    BLAST RID: {rid}")

    deadline = time.time() + timeout_total
    while time.time() < deadline:
        time.sleep(15)
        try:
            get_url = f"{BLAST_URL}?CMD=Get&RID={rid}&FORMAT_TYPE=XML"
            with urllib.request.urlopen(get_url, timeout=20) as r2:
                t2 = r2.read().decode()

            if 'Status=WAITING' in t2:
                continue
            if '<BlastOutput>' in t2:
                hits_def = re.findall(r'<Hit_def>([^<]+)</Hit_def>', t2)
                hits_e = re.findall(r'<Hsp_evalue>([^<]+)</Hsp_evalue>', t2)

                stats = {'total': 0, 'target': 0, 'offtarget': 0}
                for i in range(len(hits_def)):
                    e = float(hits_e[i]) if i < len(hits_e) else 999
                    if e >= 1:
                        continue
                    stats['total'] += 1
                    if gene_name and gene_name in hits_def[i]:
                        stats['target'] += 1
                    else:
                        stats['offtarget'] += 1
                return stats
        except:
            continue

    return {'error': 'timeout', 'total': 0, 'target': 0, 'offtarget': 0}


# ─────────────────────────────────────────────
# 第5步: 报告生成
# ─────────────────────────────────────────────
def generate_report(selected, gene_name, transcript_id, cds_len, output_path):
    """生成 .docx 报告"""
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
    except ImportError:
        print("提示: 安装 python-docx 以获得 .docx 报告: pip install python-docx")
        # 回退到txt
        txt_path = output_path.replace('.docx', '.txt')
        with open(txt_path, 'w') as f:
            f.write(f"{gene_name} siRNA设计报告\n{'='*50}\n")
            for i, c in enumerate(selected):
                f.write(f"\nsiRNA-{i+1}:\n")
                f.write(f"  Sense: 5'-{c['ss']}TT-3'\n")
                f.write(f"  Antisense: 5'-{c['as']}TT-3'\n")
                f.write(f"  DSIR Score: {c['score']}\n")
        print(f"报告已保存: {txt_path}")
        return

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    title = doc.add_heading(f'{gene_name} siRNA 设计报告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run(f'基因: ').bold = True; p.add_run(f'{gene_name}')
    p.add_run(f'    参考转录本: ').bold = True; p.add_run(f'{transcript_id} (CDS: {cds_len}bp)')

    doc.add_paragraph('设计流程: DSIR 19nt模式 + 文档筛选标准 + BLAST验证')
    doc.add_paragraph()

    for i, c in enumerate(selected):
        doc.add_heading(f'siRNA-{i+1}', level=1)
        t = doc.add_table(rows=6, cols=4)
        t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER

        blast_info = c.get('blast', {})
        if blast_info.get('error'):
            blast_str = f'BLAST: {blast_info["error"]}'
        else:
            blast_str = f'靶基因: {blast_info.get("target", "?")}  脱靶: {blast_info.get("offtarget", "?")}'

        rows_data = [
            [f'序列{i+1}', '内在稳定性', f'F1={c["f1"]} R1={c["r1"]} F2=A', ''],
            ['F', f'{c["ss"]}TT', 'BLAST结果', blast_str],
            ['R', f'{c["as"]}TT', 'GC含量', f'{c["gc_pct"]:.1f}% ({c["gc"]}/19)'],
            ['长度', '21nt (19nt+TT)', '连续碱基', '-'],
            ['位置', f'{c["aa"]}aa ({c["pos"]}bp)', '重复序列', '-'],
            ['DSIR', f'Score={c["score"]}  Corrected={c["corrected"]}', '', ''],
        ]
        for ri, rd in enumerate(rows_data):
            for j, v in enumerate(rd):
                if v: t.rows[ri].cells[j].text = v
        t.rows[0].cells[0].merge(t.rows[0].cells[1])
        t.rows[0].cells[2].merge(t.rows[0].cells[3])
        doc.add_paragraph()

    # 订购表
    doc.add_heading('订购序列', level=1)
    ot = doc.add_table(rows=3, cols=3)
    ot.style = 'Table Grid'; ot.alignment = WD_TABLE_ALIGNMENT.CENTER
    ot.rows[0].cells[0].text = ''
    for i, c in enumerate(selected):
        ot.rows[0].cells[i+1].text = f'siRNA-{i+1}'
        ot.rows[1].cells[i+1].text = f"5'-{c['ss']}TT-3'"
        ot.rows[2].cells[i+1].text = f"5'-{c['as']}TT-3'"
    ot.rows[1].cells[0].text = 'Sense (F)'
    ot.rows[2].cells[0].text = 'Antisense (R)'

    doc.add_paragraph()
    spacing = abs(selected[0]['pos'] - selected[1]['pos']) if len(selected) >= 2 else 0
    doc.add_paragraph(f'间距: {spacing}bp  |  合成时TT备注为dTdT  |  建议HPLC纯化')

    doc.save(output_path)
    print(f"报告已保存: {output_path}")


# ─────────────────────────────────────────────
# 主流程
# ─────────────────────────────────────────────
def design_sirnas(gene_name=None, species="Homo sapiens", transcript_id=None,
                  cds_seq=None, output_dir=".", verbose=True):
    """完整siRNA设计流程"""
    print("=" * 60)
    print(f"siRNA 设计流程 - {gene_name or transcript_id or '自定义序列'}")
    print("=" * 60)

    # 第1步: 获取CDS
    print("\n[步骤1] 获取CDS序列...")
    if cds_seq:
        seq = cds_seq.upper().replace(' ', '').replace('\n', '')
        transcript_info = "custom_sequence"
        cds_len = len(seq)
        print(f"  使用自定义序列 ({cds_len}bp)")
    elif transcript_id:
        seq, header = fetch_cds_from_ncbi(transcript_id)
        transcript_info = transcript_id
        cds_len = len(seq)
        print(f"  {transcript_id} → CDS: {cds_len}bp")
    elif gene_name:
        tid, seq, header = search_transcript(gene_name, species)
        transcript_info = tid
        cds_len = len(seq)
        print(f"  {gene_name} → {tid} → CDS: {cds_len}bp ({cds_len//3}aa)")
    else:
        raise ValueError("需要指定 --gene, --transcript 或 --cds")

    if not seq.startswith('ATG'):
        print("  ⚠ 序列不是以ATG开头（可能不是完整CDS）")

    # 第2步: DSIR
    print(f"\n[步骤2] DSIR预测 (19nt, threshold=90)...")
    results = submit_dsir(seq, threshold=90)
    print(f"  DSIR返回: {len(results)} 条候选 (Score≥90)")

    if len(results) == 0:
        print("  ⚠ Score≥90的候选为0，尝试降低threshold")
        for t in [80, 70, 60, 50]:
            results = submit_dsir(seq, threshold=t)
            print(f"  threshold={t}: {len(results)} 条")
            if len(results) >= 10:
                break

    # 第3步: 筛选
    print(f"\n[步骤3] 筛选...")
    print(f"  条件: 位置>{100}aa, 第2位=A, GC 8-11, 无bad patterns")

    # 先按 Corrected Score ≥ 90 筛选
    high_score = [r for r in results if r['corrected'] >= 90]
    print(f"  Corrected Score ≥ 90: {len(high_score)} 条")

    candidates = filter_candidates(high_score if len(high_score) >= 2 else results,
                                    start_aa=100, verbose=False)

    if len(candidates) < 2:
        print(f"  Corrected Score≥90筛选后候选: {len(candidates)} 条 (<2)")
        print(f"  放宽至所有DSIR结果...")
        candidates = filter_candidates(results, start_aa=100, verbose=False)

    print(f"  通过手动筛选: {len(candidates)} 条")

    if len(candidates) == 0:
        print("  ✗ 无候选通过筛选!")
        return

    selected = pick_best_two(candidates)
    print(f"\n  选定 {len(selected)} 条:")
    for i, c in enumerate(selected):
        print(f"    siRNA-{i+1}: pos={c['pos']}({c['aa']}aa), "
              f"GC={c['gc']}/19, Score={c['score']}, Corrected={c['corrected']}")

    # 第4步: BLAST
    print(f"\n[步骤4] BLAST验证...")
    for i, c in enumerate(selected):
        print(f"  siRNA-{i+1} ({c['ss']})...")
        result = blast_sequence(c['ss'], gene_name=gene_name)
        c['blast'] = result
        if 'error' in result:
            print(f"    ⚠ {result['error']}")
        else:
            print(f"    总命中: {result['total']}, 靶基因: {result['target']}, "
                  f"人类脱靶: {result['offtarget']}")

    # 第5步: 报告
    print(f"\n[步骤5] 生成报告...")
    output_path = os.path.join(output_dir, f'{gene_name or "siRNA"}_siRNA设计报告.docx')
    generate_report(selected, gene_name or "Target", transcript_info, cds_len, output_path)

    # 终端输出摘要
    print(f"\n{'='*60}")
    print("设计完成!")
    print("="*60)
    for i, c in enumerate(selected):
        print(f"\n  siRNA-{i+1}  订购序列:")
        print(f"  F: 5'-{c['ss']}TT-3'")
        print(f"  R: 5'-{c['as']}TT-3'")
        print(f"  DSIR Score: {c['score']} | 位置: {c['aa']}aa | GC: {c['gc_pct']:.1f}%")

    if len(selected) >= 2:
        print(f"\n  间距: {abs(selected[0]['pos'] - selected[1]['pos'])}bp")

    return selected


# ─────────────────────────────────────────────
# CLI入口
# ─────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description='siRNA设计工具 - 按文档流程自动设计siRNA',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  python sirna_design.py --gene SUN1
  python sirna_design.py --gene SUN1 --species "Homo sapiens" --output ./results
  python sirna_design.py --transcript NM_001130965.3
  python sirna_design.py --cds my_sequence.fasta
        """
    )
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument('--gene', help='基因名 (如 SUN1)')
    group.add_argument('--transcript', help='转录本ID (如 NM_001130965.3)')
    group.add_argument('--cds', help='FASTA格式CDS文件路径')

    parser.add_argument('--species', default='Homo sapiens', help='物种（默认: Homo sapiens）')
    parser.add_argument('--output', '-o', default='.', help='输出目录')
    parser.add_argument('--quiet', '-q', action='store_true', help='安静模式')

    args = parser.parse_args()

    if args.cds:
        seq, header = read_fasta(args.cds)
        design_sirnas(cds_seq=seq, output_dir=args.output, verbose=not args.quiet)
    else:
        design_sirnas(gene_name=args.gene, species=args.species,
                      transcript_id=args.transcript, output_dir=args.output,
                      verbose=not args.quiet)


if __name__ == '__main__':
    main()
